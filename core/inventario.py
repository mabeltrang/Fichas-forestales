"""Lectura del inventario forestal (formato Unergy) y cálculo de variables."""
from __future__ import annotations

import math
import re
import unicodedata
from io import BytesIO

import openpyxl
import pandas as pd

# Nombre interno -> posibles encabezados en el Excel (normalizados)
COLUMNAS = {
    "id": ["id"],
    "cobertura": ["cobertura"],
    "nombre_comun": ["nombre comun"],
    "nombre_cientifico": ["nombre cientifico", "especie"],
    "familia": ["familia"],
    "cap_a": ["cap a (cm)", "cap (cm)"],
    "dap_a": ["dap a (m)", "dap (m)"],
    "cap_b": ["cap b (cm)"],
    "dap_b": ["dap b (m)"],
    "cap_c": ["cap c (cm)"],
    "dap_c": ["dap c (m)"],
    "ht": ["ht (m)"],
    "hc": ["hc (m)"],
    "dc_ns": ["dc (n-s) (m)"],
    "dc_eo": ["dc (e-o) (m)"],
    "x": ["coordenada x", "longitud", "este"],
    "y": ["coordenada y", "latitud", "norte"],
    "municipio": ["municipio"],
    "departamento": ["departamento"],
    "estado": ["estado fitosanitario"],
    "motivo": ["motivo de intervencion"],
    "observaciones": ["observaciones"],
    "medida": ["medida de manejo"],
}

ENCABEZADO = ["PROYECTO", "PROPIETARIO", "UBICACION", "FECHA"]


def norm(texto) -> str:
    if texto is None:
        return ""
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def _num(v):
    """Devuelve float o None. Las fórmulas sin valor calculado se tratan como vacías."""
    if v is None or (isinstance(v, str) and (v.strip() == "" or v.startswith("="))):
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def _txt(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.startswith("=") else s


def leer_inventario(archivo: bytes | BytesIO) -> tuple[pd.DataFrame, dict]:
    """Lee la hoja de inventario. Devuelve (tabla, metadatos del encabezado)."""
    data = archivo if isinstance(archivo, BytesIO) else BytesIO(archivo)
    wb = openpyxl.load_workbook(data, data_only=False)
    ws = wb.active

    # Fila de encabezados = primera fila cuya primera celda dice "ID"
    fila_hdr = None
    for r in range(1, min(ws.max_row, 30) + 1):
        if norm(ws.cell(r, 1).value) == "id":
            fila_hdr = r
            break
    if fila_hdr is None:
        raise ValueError("No encontré la fila de encabezados (una celda 'ID' en la columna A).")

    # Metadatos del bloque superior (etiqueta y valor en la misma fila)
    meta = {}
    for r in range(1, fila_hdr):
        celdas = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for i, v in enumerate(celdas):
            if norm(v).upper() in ENCABEZADO:
                valor = next((x for x in celdas[i + 1:] if x not in (None, "")), "")
                meta[norm(v)] = _txt(valor)

    hdr = {norm(ws.cell(fila_hdr, c).value): c for c in range(1, ws.max_column + 1)}
    idx = {}
    for clave, opciones in COLUMNAS.items():
        for o in opciones:
            if o in hdr:
                idx[clave] = hdr[o]
                break

    registros = []
    for r in range(fila_hdr + 1, ws.max_row + 1):
        val = {k: ws.cell(r, c).value for k, c in idx.items()}
        if not _txt(val.get("id")):
            continue
        registros.append(_procesar_fila(val))

    return pd.DataFrame(registros), meta


def _procesar_fila(v: dict) -> dict:
    # DAP por fuste: usa el DAP si viene numérico; si no, lo calcula desde CAP
    daps = []
    for s in ("a", "b", "c"):
        dap = _num(v.get(f"dap_{s}"))
        cap = _num(v.get(f"cap_{s}"))
        if dap is None and cap is not None:
            dap = (cap / 100) / math.pi
        if dap:
            daps.append(dap)
    # Multifustal: DAP equivalente = raíz de la suma de cuadrados (misma área basal)
    dap_eq = math.sqrt(sum(d * d for d in daps)) if daps else None

    nombre_inv = _txt(v.get("nombre_comun"))
    cientifico_inv = _txt(v.get("nombre_cientifico"))

    return {
        "ID": _txt(v.get("id")),
        "nombre_inventario": nombre_inv or cientifico_inv,
        "cientifico_inventario": cientifico_inv,
        "familia_inventario": _txt(v.get("familia")),
        "n_fustes": len(daps),
        "DAP (m)": round(dap_eq, 3) if dap_eq else None,
        "HT (m)": _num(v.get("ht")),
        "DC N-S (m)": _num(v.get("dc_ns")),
        "DC E-O (m)": _num(v.get("dc_eo")),
        "X": _num(v.get("x")),
        "Y": _num(v.get("y")),
        "Municipio": _txt(v.get("municipio")),
        "Departamento": _txt(v.get("departamento")),
        "Estado fitosanitario": _txt(v.get("estado")),
        "Motivo": _txt(v.get("motivo")),
        "Observaciones inventario": _txt(v.get("observaciones")),
        "Medida inventario": _txt(v.get("medida")),
    }


def orden_natural(id_: str):
    partes = re.split(r"(\d+)", str(id_))
    return [int(p) if p.isdigit() else p.lower() for p in partes]
