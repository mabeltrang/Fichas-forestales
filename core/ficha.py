"""Llena la plantilla de ficha (tabla única) y arma un .docx con una ficha por página."""
from __future__ import annotations

import copy
import math
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from PIL import Image

from core.inventario import norm

CARACTERES_POR_RENGLON = 95  # observación, letra de la plantilla, ancho completo de la tabla
ALTO_RENGLON_CM = 0.5
FOTO_ALTO_MIN_CM = 2.5

# Texto del marcador en la plantilla (normalizado, sin comillas) -> clave del dict de la ficha
MARCADORES = {
    "proyecto": "proyecto",
    "municipio, departamento": "ubicacion",
    "coordenadas": "coordenadas",
    "nombre": "nombre",
    "especie": "especie",
    "familia": "familia",
    "ht": "ht",
    "dap": "dap",
    "bueno/regular/ malo": "estado",
    "%": "pct_copa",
    "si - no": "fauna",
    "¿nidos / colmenas / epifitas, etc?": "tipo_fauna",
    "nidos / colmenas / epifitas, etc?": "tipo_fauna",
    "medida de manejo": "medida",
}
# Marcadores que en la plantilla no van entre comillas
SIN_COMILLAS = {"bueno/regular/ malo", "si - no", "¿nidos / colmenas / epifitas, etc?", "nidos / colmenas / epifitas, etc?"}
ETIQUETAS = {  # se conserva la etiqueta y se agrega el valor
    "elaboro:": ("ELABORÓ: ", "elaboro"),
    "matricula profesional:": ("MATRÍCULA PROFESIONAL: ", "matricula"),
    "fecha:": ("FECHA: ", "fecha"),
}


def _texto(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def _clave(texto: str) -> str:
    return norm(texto.replace("“", "").replace("”", "").replace('"', ""))


def _poner_texto(p, texto: str) -> None:
    """Reemplaza el texto del párrafo conservando el formato del primer run con texto."""
    runs = [r for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None]
    if not runs:
        r = p.makeelement(qn("w:r"), {})
        p.append(r)
        runs = [r]
        r.append(r.makeelement(qn("w:t"), {}))
    primero = runs[0]
    ts = primero.findall(qn("w:t"))
    ts[0].text = texto
    ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for t in ts[1:]:
        primero.remove(t)
    for r in runs[1:]:
        p.remove(r)


def _vaciar(p) -> None:
    for r in p.findall(qn("w:r")):
        p.remove(r)


# ---------- Nombres científicos en la observación ----------

_SP = {"sp", "sp.", "spp", "spp."}


def _patrones_nombre(nombre: str) -> list[tuple[re.Pattern, tuple | None]]:
    """Patrones que reconocen un nombre científico en texto libre.
    Devuelve (patrón, (género, epíteto)) para binomiales y (patrón, None) para 'Género sp.'."""
    t = (nombre or "").replace("\xa0", " ").split()
    if not t or not t[0][:1].isalpha():
        return []
    gen = t[0].strip("().,;")
    if not gen[:1].isupper() or len(gen) < 3:
        return []
    if len(t) >= 2 and t[1].lower() not in _SP and re.fullmatch(r"[a-zë-]+", t[1].strip("().,;")):
        ep = t[1].strip("().,;")
        par = (gen, ep)
        return [(re.compile(rf"\b{re.escape(gen)}\s+{re.escape(ep)}\b", re.I), par),
                (re.compile(rf"\b{re.escape(gen[0])}\.\s*{re.escape(ep)}\b", re.I), par)]
    # Solo género: va en cursiva si le sigue sp./spp. o cierra paréntesis (evita el nombre común "Acacia")
    return [(re.compile(rf"\b{re.escape(gen)}\b(?=\s+spp?\b|\s*\))"), None)]


def compilar_nombres(nombres) -> list[tuple[re.Pattern, tuple | None]]:
    patrones, vistos = [], set()
    for n in nombres:
        clave = " ".join(str(n or "").split()[:2]).lower()
        if clave and clave not in vistos:
            vistos.add(clave)
            patrones += _patrones_nombre(str(n))
    return patrones


def segmentar_observacion(texto: str, patrones, ya_citadas: set) -> list[tuple[str, bool]]:
    """Parte el texto en (fragmento, cursiva). La primera vez que aparece una especie en el
    documento va con el género completo; de ahí en adelante 'H. lusitanica'."""
    hallazgos = []
    for pat, par in patrones:
        for m in pat.finditer(texto):
            hallazgos.append((m.start(), -(m.end() - m.start()), m.end(), par))
    hallazgos.sort()
    segs, pos = [], 0
    for ini, _, fin, par in hallazgos:
        if ini < pos:  # se solapa con uno ya tomado
            continue
        if ini > pos:
            segs.append((texto[pos:ini], False))
        if par is None:
            segs.append((texto[ini:fin], True))
        else:
            gen, ep = par
            clave = (gen.lower(), ep.lower())
            segs.append((f"{gen[0].upper()}. {ep}" if clave in ya_citadas else f"{gen} {ep}", True))
            ya_citadas.add(clave)
        pos = fin
    if pos < len(texto):
        segs.append((texto[pos:], False))
    return segs


def _poner_segmentos(p, parent, segs: list[tuple[str, bool]]) -> None:
    """Como _poner_texto, pero con fragmentos en cursiva (mismo formato base de la plantilla)."""
    _poner_texto(p, "")
    base = next(r for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None)
    par = Paragraph(p, parent)
    for texto, cursiva in segs:
        r = copy.deepcopy(base)
        t = r.find(qn("w:t"))
        t.text = texto
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        base.addprevious(r)
        if cursiva:
            Run(r, par).italic = True
    p.remove(base)


# ---------- Tabla ----------

def _fusionar_siluetas(tbl) -> None:
    """Con una sola foto, las filas SILUETA y FOTO quedan con una celda a todo el ancho."""
    for tr in tbl.findall(qn("w:tr")):
        txt = _clave(_texto(tr))
        if not (txt.startswith("silueta") or "foto 1" in txt):
            continue
        celdas = tr.findall(qn("w:tc"))
        if len(celdas) < 2:
            continue
        span = ancho = 0
        for tc in celdas:
            tcpr = tc.find(qn("w:tcPr"))
            gs = tcpr.find(qn("w:gridSpan")) if tcpr is not None else None
            span += int(gs.get(qn("w:val"))) if gs is not None else 1
            w = tcpr.find(qn("w:tcW")) if tcpr is not None else None
            if w is not None and (w.get(qn("w:type")) in (None, "dxa")):
                ancho += int(w.get(qn("w:w"), "0"))
        primera = celdas[0]
        for tc in celdas[1:]:
            tr.remove(tc)
        tcpr = primera.find(qn("w:tcPr"))
        gs = tcpr.find(qn("w:gridSpan"))
        if gs is None:
            gs = tcpr.makeelement(qn("w:gridSpan"), {})
            w = tcpr.find(qn("w:tcW"))
            if w is not None:
                w.addnext(gs)
            else:
                tcpr.insert(0, gs)
        gs.set(qn("w:val"), str(span))
        w = tcpr.find(qn("w:tcW"))
        if w is not None and ancho:
            w.set(qn("w:w"), str(ancho))
            w.set(qn("w:type"), "dxa")
        if txt.startswith("silueta"):
            for p in primera.iter(qn("w:p")):
                if _texto(p).strip():
                    _poner_texto(p, "SILUETA")
                    break


def _ajustar_filas(tbl, alto_foto_twips: int) -> None:
    """La plantilla trae alturas 'exactas' que recortan texto: se pasan a 'mínimo'.
    La fila de fotos se fija a la altura de la imagen y se elimina la fila vacía de abajo."""
    filas = tbl.findall(qn("w:tr"))
    for i, tr in enumerate(filas[:-1]):
        if "foto 1" in _clave(_texto(tr)) and not _clave(_texto(filas[i + 1])):
            tbl.remove(filas[i + 1])
            for vm in tr.iter(qn("w:vMerge")):
                vm.getparent().remove(vm)
            break
    for tr in tbl.findall(qn("w:tr")):
        txt = _clave(_texto(tr))
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            continue
        h = trpr.find(qn("w:trHeight"))
        if h is None:
            continue
        if "foto 1" in txt:
            h.set(qn("w:hRule"), "exact")
            h.set(qn("w:val"), str(alto_foto_twips))
        elif tbl.findall(qn("w:tr")).index(tr) == 0:
            continue  # encabezado con logo: se deja igual
        else:
            h.set(qn("w:hRule"), "atLeast")


def _insertar_foto(p_el, parent, foto: BytesIO | None, ancho_cm: float, alto_cm: float) -> None:
    _vaciar(p_el)
    ppr = p_el.find(qn("w:pPr"))
    if ppr is not None:  # el interlineado exacto de la plantilla recorta la imagen
        for sp in ppr.findall(qn("w:spacing")):
            ppr.remove(sp)
        for ind in ppr.findall(qn("w:ind")):
            ppr.remove(ind)
    par = Paragraph(p_el, parent)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if foto is None:
        par.add_run("Sin registro fotográfico")
        return
    w, h = Image.open(foto).size
    foto.seek(0)
    escala = min(ancho_cm / w, alto_cm / h)
    par.add_run().add_picture(foto, width=Cm(w * escala), height=Cm(h * escala))


def _salto_pagina(body_parent):
    p = body_parent.makeelement(qn("w:p"), {})
    ppr = p.makeelement(qn("w:pPr"), {})
    pb = ppr.makeelement(qn("w:pageBreakBefore"), {})
    sp = ppr.makeelement(qn("w:spacing"), {qn("w:before"): "0", qn("w:after"): "0", qn("w:line"): "20", qn("w:lineRule"): "exact"})
    ppr.append(pb)
    ppr.append(sp)
    p.append(ppr)
    return p


def _parrafo_minimo(body_parent):
    p = _salto_pagina(body_parent)
    p.find(qn("w:pPr")).remove(p.find(qn("w:pPr")).find(qn("w:pageBreakBefore")))
    return p


def _llenar(tbl, f: dict, parent, foto_ancho_cm: float, foto_alto_cm: float,
            obs_segs: list[tuple[str, bool]] | None = None) -> None:
    n_id = n_diam = n_foto = 0
    for p in list(tbl.iter(qn("w:p"))):
        texto = _texto(p).strip()
        if not texto:
            continue
        k = _clave(texto)

        if k.replace(" ", "").startswith("fichano"):  # la plantilla dice "FICHA No.con"
            _poner_texto(p, "FICHA No.")
        elif "intervenir" in k:
            ppr = p.find(qn("w:pPr"))
            if ppr is not None:  # la plantilla trae sangría francesa: se quita y se centra
                for ind in ppr.findall(qn("w:ind")):
                    ppr.remove(ind)
            Paragraph(p, parent).alignment = WD_ALIGN_PARAGRAPH.CENTER
            _poner_texto(p, f"IDENTIFICACIÓN DE ÁRBOLES A INTERVENIR - {f['titulo']}")
        elif k == "id":
            _poner_texto(p, str(f["ficha_no"] if n_id == 0 else f["id"]))
            n_id += 1
        elif k == "diametro":
            _poner_texto(p, f["dc_ns"] if n_diam == 0 else f["dc_eo"])
            n_diam += 1
        elif k in ("foto 1", "foto 2"):
            if f.get("_foto_unica"):
                _insertar_foto(p, parent, f["foto1"] or f["foto2"], foto_ancho_cm * 2, foto_alto_cm)
            else:
                foto = f["foto1"] if k == "foto 1" else f["foto2"]
                _insertar_foto(p, parent, foto, foto_ancho_cm, foto_alto_cm)
            n_foto += 1
        elif k.startswith("arbol de la especie"):
            ppr = p.find(qn("w:pPr"))
            if ppr is not None:  # sin sangría para ganar ancho de renglón
                for ind in ppr.findall(qn("w:ind")):
                    ppr.remove(ind)
            if obs_segs:
                _poner_segmentos(p, parent, obs_segs)
            else:
                _poner_texto(p, f["observacion"])
        elif k in ETIQUETAS:
            etiqueta, campo = ETIQUETAS[k]
            _poner_texto(p, etiqueta + str(f[campo]))
        elif k in MARCADORES and (k in SIN_COMILLAS or "“" in texto or '"' in texto):
            _poner_texto(p, str(f[MARCADORES[k]]))


def generar_documento(plantilla: bytes, fichas: list[dict], foto_ancho_cm=7.4, foto_alto_cm=6.2,
                      otros_nombres=()) -> bytes:
    """otros_nombres: nombres científicos adicionales (p. ej. el catálogo) para poner en cursiva
    si aparecen en las observaciones."""
    doc = Document(BytesIO(plantilla))
    body = doc.element.body
    tbl_base = body.find(qn("w:tbl"))
    if tbl_base is None:
        raise ValueError("La plantilla no tiene la tabla de la ficha.")
    sectpr = body.find(qn("w:sectPr"))
    for el in list(body):
        if el is not sectpr:
            body.remove(el)

    patrones = compilar_nombres([f.get("especie") for f in fichas] + list(otros_nombres))
    ya_citadas: set = set()  # especies ya escritas completas en este documento

    for i, f in enumerate(fichas):
        if i > 0:
            sectpr.addprevious(_salto_pagina(body))
        # Cada renglón de observación por encima de 2 se descuenta del alto de las fotos
        renglones = math.ceil(len(f["observacion"]) / CARACTERES_POR_RENGLON)
        alto = max(FOTO_ALTO_MIN_CM, foto_alto_cm - max(0, renglones - 2) * ALTO_RENGLON_CM)
        tbl = copy.deepcopy(tbl_base)
        _ajustar_filas(tbl, int((alto + 0.25) * 567))
        f = dict(f, _foto_unica=(f.get("foto1") is None) or (f.get("foto2") is None))
        if f["_foto_unica"]:
            _fusionar_siluetas(tbl)
        sectpr.addprevious(tbl)
        segs = segmentar_observacion(f["observacion"], patrones, ya_citadas)
        _llenar(tbl, f, doc._body, foto_ancho_cm, alto, segs)
    sectpr.addprevious(_parrafo_minimo(body))  # Word exige un párrafo después de una tabla

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
